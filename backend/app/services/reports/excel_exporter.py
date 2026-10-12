"""Excel Report Exporter using openpyxl (FR-063, Spec §49.11)."""
from __future__ import annotations

import io
from typing import Any, Dict, List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table as XLTable, TableStyleInfo


def _str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (list, tuple)):
        return ", ".join(str(v) for v in value)
    if isinstance(value, dict):
        return "; ".join(f"{k}: {v}" for k, v in value.items())
    return str(value)


def _pct(value: Any) -> str:
    if value is None or value == "":
        return "—"
    try:
        return f"{float(value)}%"
    except (TypeError, ValueError):
        return _str(value)


def _style_header(cell, header_fill, header_font, thin_border):
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border


def _style_row(cell, regular_font, thin_border, wrap=True, bold=False, font=None):
    cell.font = font or regular_font
    cell.border = thin_border
    cell.alignment = Alignment(horizontal="left" if cell.column > 1 or bold else "left", vertical="center", wrap_text=wrap)


def _autosize(sheet, max_width=52, min_width=12):
    for col in sheet.columns:
        try:
            col_letter = get_column_letter(col[0].column)
        except Exception:
            continue
        max_len = 0
        for cell in col:
            val_str = _str(cell.value or "")
            # Look at the first line only for width estimation.
            first = val_str.split("\n")[0]
            if len(first) > max_len and len(first) < 120:
                max_len = len(first)
        sheet.column_dimensions[col_letter].width = min(max(max_len + 4, min_width), max_width)


def export_report_to_excel(report: Dict[str, Any]) -> bytes:
    """Render a structured reading report into an in-memory XLSX workbook."""
    wb = openpyxl.Workbook()

    header_fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
    sub_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    gold_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="312E81")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    section_font = Font(name="Calibri", size=12, bold=True, color="1E293B")
    bold_font = Font(name="Calibri", size=10, bold=True, color="1E293B")
    regular_font = Font(name="Calibri", size=10, color="334155")
    small_font = Font(name="Calibri", size=9, color="64748B")
    italic_font = Font(name="Calibri", size=9, italic=True, color="64748B")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    wrap_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    wrap_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # SHEET 1: OVERVIEW
    ws = wb.active
    ws.title = "Overview"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.orientation = "portrait"
    ws.print_title_rows = "1:4"
    ws.freeze_panes = "A5"

    ws["A1"] = "MysticLens Spiritual Intelligence"
    ws["A1"].font = italic_font
    ws["A2"] = _str(report.get("title", "Reading Report"))
    ws["A2"].font = title_font
    ws["A3"] = f"{_str(report.get('report_type', '')).title()}  •  {_str(report.get('user_name', ''))}  •  {_str(report.get('created_at', ''))}"
    ws["A3"].font = subtitle_font

    ws["A4"] = "REPORT OVERVIEW"
    ws["A4"].font = section_font

    confidence = report.get("confidence", {}) or {}
    metadata_rows = [
        ("Report ID", _str(report.get("report_id", ""))),
        ("Reading ID", _str(report.get("reading_id", ""))),
        ("Report Type", _str(report.get("report_type", "")).title()),
        ("Seeker", _str(report.get("user_name", ""))),
        ("Created At", _str(report.get("created_at", ""))),
        ("Reading Status", _str(report.get("reading_status", "completed") or "completed")),
        ("Review State", "Reviewed ✓" if report.get("reviewed") else "Awaiting review"),
        ("Confidence Score", _pct(confidence.get("score", 0))),
        ("Confidence Tier", _str(confidence.get("label", ""))),
    ]
    if report.get("report_type") == "tarot":
        metadata_rows.insert(5, ("Sacred Question", _str(report.get("question", ""))))
        spread = report.get("spread_details", {}) or {}
        metadata_rows.insert(6, ("Spread", f"{_str(spread.get('spread_name', ''))} ({_str(spread.get('total_cards', 0))} cards)"))
    else:
        palm = report.get("palm_features", {}) or {}
        metadata_rows.insert(5, ("Hand Element", f"{_str(palm.get('element', ''))} Hand — {_str(palm.get('hand_detected', ''))} hand")),
        metadata_rows.insert(6, ("Hand Traits", _str(palm.get("traits", [])))),

    for idx, (label, val) in enumerate(metadata_rows, start=5):
        a = ws[f"A{idx}"]
        b = ws[f"B{idx}"]
        a.value = label
        a.font = bold_font
        a.border = thin_border
        a.fill = zebra_fill if idx % 2 == 0 else PatternFill(fill_type=None)
        a.alignment = Alignment(horizontal="left", vertical="center")
        b.value = val
        b.font = regular_font
        b.border = thin_border
        b.alignment = wrap_left
        ws.row_dimensions[idx].height = 18 if len(val) < 80 else 30

    summary_row = len(metadata_rows) + 6
    ws[f"A{summary_row}"] = "EXECUTIVE SUMMARY"
    ws[f"A{summary_row}"].font = section_font
    ws[f"A{summary_row+1}"] = _str(report.get("summary", ""))
    ws[f"A{summary_row+1}"].font = regular_font
    ws[f"A{summary_row+1}"].alignment = wrap_left
    ws.row_dimensions[summary_row + 1].height = 45
    ws.merge_cells(f"A{summary_row+1}:B{summary_row+1}")

    if report.get("question") and report.get("report_type") == "tarot":
        ws[f"A{summary_row+2}"] = f"Inquiry: “{_str(report.get('question'))}”"
        ws[f"A{summary_row+2}"].font = italic_font
        ws.merge_cells(f"A{summary_row+2}:B{summary_row+2}")

    disclaimer_row = summary_row + 4
    ws[f"A{disclaimer_row}"] = "ETHICS & SAFETY NOTICE (§49.19)"
    ws[f"A{disclaimer_row}"].font = bold_font
    ws[f"A{disclaimer_row}"].fill = gold_fill
    ws[f"A{disclaimer_row+1}"] = _str(report.get("disclaimer", ""))
    ws[f"A{disclaimer_row+1}"].font = italic_font
    ws[f"A{disclaimer_row+1}"].alignment = wrap_left
    ws.row_dimensions[disclaimer_row + 1].height = 40
    ws.merge_cells(f"A{disclaimer_row+1}:B{disclaimer_row+1}")

    if report.get("reading_notes"):
        ws[f"A{disclaimer_row+3}"] = "SEEKER JOURNAL NOTES"
        ws[f"A{disclaimer_row+3}"].font = bold_font
        ws[f"A{disclaimer_row+4}"] = _str(report.get("reading_notes"))
        ws[f"A{disclaimer_row+4}"].font = italic_font
        ws[f"A{disclaimer_row+4}"].alignment = wrap_left
        ws.merge_cells(f"A{disclaimer_row+4}:B{disclaimer_row+4}")

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 90

    # SHEET 2: ANALYSIS DETAILS
    ws_details = wb.create_sheet(title="Analysis Details")
    ws_details.sheet_properties.pageSetUpPr.fitToPage = True
    ws_details.page_setup.fitToWidth = 1
    ws_details.page_setup.fitToHeight = 0
    ws_details.page_setup.orientation = "landscape"
    ws_details.freeze_panes = "A3"

    report_type = report.get("report_type")

    def _write_table(start_row: int, title: str, headers: List[str], rows: List[List[Any]], widths: List[int] | None = None):
        ws_details.cell(row=start_row, column=1, value=title).font = section_font
        header_row = start_row + 1
        for col_idx, h in enumerate(headers, start=1):
            cell = ws_details.cell(row=header_row, column=col_idx, value=h)
            _style_header(cell, header_fill, header_font, thin_border)
        ws_details.row_dimensions[header_row].height = 28
        current = header_row + 1
        for r_idx, row_vals in enumerate(rows):
            max_lines = 1
            for col_idx, val in enumerate(row_vals, start=1):
                cell = ws_details.cell(row=current, column=col_idx, value=val)
                cell.font = regular_font
                cell.border = thin_border
                cell.alignment = wrap_left if col_idx != 1 else Alignment(horizontal="left", vertical="center", wrap_text=True)
                max_lines = max(max_lines, max(1, len(str(val or "")) // 60 + 1))
            if r_idx % 2 == 0:
                for col_idx in range(1, len(headers) + 1):
                    ws_details.cell(row=current, column=col_idx).fill = zebra_fill
            ws_details.row_dimensions[current].height = min(max(22, max_lines * 15), 90)
            current += 1
        # Excel Table for filtering
        if rows:
            ref = f"{get_column_letter(1)}{header_row}:{get_column_letter(len(headers))}{current - 1}"
            tab = XLTable(displayName=f"Table{abs(hash(title)) % 9000 + 1000}", ref=ref)
            tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
            try:
                ws_details.add_table(tab)
            except Exception:
                pass
            ws_details.auto_filter.ref = ref
        return current

    if report_type == "palmistry":
        palm = report.get("palm_features", {}) or {}
        finger = report.get("finger_structure", {}) or {}
        ws_details["A1"] = (
            f"PALMISTRY BIOMETRICS — {_str(palm.get('element', ''))} Hand  •  {_str(palm.get('hand_detected', ''))} hand  •  "
            f"{_str(palm.get('landmarks_count', 21))} landmarks  •  Confidence {_pct((report.get('confidence', {}) or {}).get('score', 0))}"
        )
        ws_details["A1"].font = section_font

        line_headers = ["Line Name", "Length", "Clarity", "Curve", "Continuity", "Detected", "Confidence %", "Interpretation"]
        line_rows = []
        for line in report.get("lines", []) or []:
            line_rows.append([
                _str(line.get("name", "")),
                _str(line.get("length", "")),
                _str(line.get("clarity", "")),
                _str(line.get("curve", "")),
                _str(line.get("continuity", "")),
                "Yes" if line.get("detected", True) else "No — faint",
                _pct(line.get("confidence", 0)),
                _str(line.get("interpretation", "")),
            ])
        current_row = _write_table(2, "PRIMARY PALM LINES", line_headers, line_rows) if line_rows else 4

        current_row += 1
        mount_headers = ["Mount Name", "Prominence", "Score %", "Element / Ruler", "Archetypal Meaning + Attributes"]
        mount_rows = []
        for mount in report.get("mounts", []) or []:
            ruler = " / ".join([x for x in [_str(mount.get("element", "")), _str(mount.get("planetary_ruler", ""))] if x])
            meaning = _str(mount.get("meaning", ""))
            attrs = mount.get("core_attributes") or []
            if attrs:
                meaning = f"{meaning}  [{', '.join(str(a) for a in attrs)}]"
            mount_rows.append([
                _str(mount.get("name", "")),
                _str(mount.get("prominence", "")),
                _pct(mount.get("score", 0)),
                ruler or "—",
                meaning,
            ])
        if mount_rows:
            current_row = _write_table(current_row, "PLANETARY MOUNTS PROMINENCE", mount_headers, mount_rows)

        current_row += 1
        ws_details.cell(row=current_row, column=1, value="HAND & FINGER METRICS").font = section_font
        current_row += 1
        metric_headers = ["Metric", "Value", "Notes"]
        metric_rows = [
            ["Hand element", _str(palm.get("element", "")), _str(palm.get("shape_description", ""))[:160]],
            ["Traits", _str(palm.get("traits", [])), "Elemental temperament markers"],
            ["Palm length ÷ width", _str(palm.get("palm_length_to_width", "")), "Biometric proportion"],
            ["Finger ÷ palm ratio", _str(palm.get("finger_to_palm_ratio", "")), "Biometric proportion"],
            ["2D:4D digit ratio", _str(finger.get("digit_ratio_2d_4d", "")), _str(finger.get("digit_ratio_interpretation", ""))[:160]],
            ["Thumb", f"{_str(finger.get('thumb_angle_degrees', ''))}° / {_str(finger.get('thumb_flexibility', ''))}", _str(finger.get("pinky_trait", ""))[:160]],
        ]
        current_row = _write_table(current_row, "", metric_headers, metric_rows)
    else:
        spread = report.get("spread_details", {}) or {}
        ws_details["A1"] = (
            f"TAROT SYMBOLICS — {_str(spread.get('spread_name', ''))}  •  {_str(spread.get('total_cards', 0))} cards  •  "
            f"Deck {_str(spread.get('deck_id', ''))}  •  Q: {_str(report.get('question', ''))[:80]}"
        )
        ws_details["A1"].font = section_font

        tarot_headers = ["#", "Position", "Position Meaning", "Card", "Arcana / Suit", "Orientation", "Keywords", "Interpretation"]
        tarot_rows = []
        for card in report.get("cards", []) or []:
            arcana_suit = _str(card.get("arcana", ""))
            if card.get("suit") and card.get("suit") != "None":
                arcana_suit += f" / {_str(card.get('suit'))}"
            tarot_rows.append([
                card.get("index", 1),
                _str(card.get("position_name", "")),
                _str(card.get("position_meaning", "")),
                _str(card.get("name", "")),
                arcana_suit,
                _str(card.get("orientation", "")),
                _str((card.get("keywords") or [])[:6]),
                _str(card.get("interpretation", "")),
            ])
        current_row = _write_table(2, "CARDS DRAWN & SPREAD POSITIONS", tarot_headers, tarot_rows) if tarot_rows else 4

        current_row += 1
        ws_details.cell(row=current_row, column=1, value="DOMAIN GUIDANCE & DEEPER SYMBOLISM").font = section_font
        current_row += 1
        deep_headers = ["Card", "Domain", "Guidance Text"]
        deep_rows = []
        for card in report.get("cards", []) or []:
            for g in (card.get("domain_guidance") or [])[:3]:
                deep_rows.append([_str(card.get("name", "")), _str(g.get("label") or g.get("domain", "")), _str(g.get("text", ""))])
            extras = (card.get("expanded_meanings") or [])[:3]
            if extras:
                deep_rows.append([_str(card.get("name", "")), "Shadow / Light insights", "; ".join(str(x) for x in extras)])
            fortune = (card.get("fortune_telling") or [])[:2]
            if fortune:
                deep_rows.append([_str(card.get("name", "")), "Traditional note", "; ".join(str(x) for x in fortune)])
        if deep_rows:
            current_row = _write_table(current_row, "", deep_headers, deep_rows)

        current_row += 1
        ws_details.cell(row=current_row, column=1, value="SPREAD SYNTHESIS").font = section_font
        current_row += 1
        synthesis = report.get("synthesis", {}) or {}
        narrative = report.get("narrative", {}) or {}
        synth_headers = ["Dimension", "Value", "Narrative"]
        arc_dist = synthesis.get("arcana_distribution", {}) or {}
        synth_rows = [
            ["Dominant element", _str(synthesis.get("elemental_dominance", "Balanced")), _str(synthesis.get("elemental_narrative", ""))[:220]],
            ["Arcana balance", f"Major {_str((arc_dist or {}).get('major', 0))} · Minor {_str((arc_dist or {}).get('minor', 0))}", _str(synthesis.get("arcana_narrative", ""))[:220]],
            ["Reversals", _str(synthesis.get("reversals_count", 0)), _str(synthesis.get("reversal_narrative", ""))[:220]],
            ["Numerology", _str(synthesis.get("numerology_profile", ""))[:80], _str(synthesis.get("numerology_narrative", ""))[:220]],
            ["Theme", _str(narrative.get("theme", "")), _str(narrative.get("summary", ""))[:220]],
        ]
        current_row = _write_table(current_row, "", synth_headers, synth_rows)

    _autosize(ws_details)

    # SHEET 3: GUIDANCE & ACTIONS
    ws_guidance = wb.create_sheet(title="Guidance & Actions")
    ws_guidance.sheet_properties.pageSetUpPr.fitToPage = True
    ws_guidance.page_setup.fitToWidth = 1
    ws_guidance.page_setup.fitToHeight = 1
    ws_guidance.freeze_panes = "A3"

    ws_guidance["A1"] = "ACTIONABLE SPIRITUAL & PRACTICAL GUIDANCE"
    ws_guidance["A1"].font = section_font

    g_headers = ["#", "Recommended Action / Reflection Step"]
    for col_idx, h in enumerate(g_headers, start=1):
        cell = ws_guidance.cell(row=2, column=col_idx, value=h)
        _style_header(cell, header_fill, header_font, thin_border)
    ws_guidance.row_dimensions[2].height = 24

    current_row = 3
    for idx, item in enumerate(report.get("actionable_guidance", []) or [], start=1):
        c1 = ws_guidance.cell(row=current_row, column=1, value=idx)
        c1.font = bold_font
        c1.border = thin_border
        c1.alignment = wrap_center
        c2 = ws_guidance.cell(row=current_row, column=2, value=_str(item))
        c2.font = regular_font
        c2.border = thin_border
        c2.alignment = wrap_left
        if current_row % 2 == 0:
            c1.fill = zebra_fill
            c2.fill = zebra_fill
        ws_guidance.row_dimensions[current_row].height = min(max(26, len(_str(item)) // 70 * 15 + 26), 90)
        current_row += 1

    # Narrative body (tarot) as follow-on block
    narrative = report.get("narrative", {}) or {}
    body_lines = narrative.get("body") or []
    if isinstance(body_lines, list) and body_lines:
        ws_guidance.cell(row=current_row + 1, column=1, value="NARRATIVE JOURNEY").font = section_font
        current_row += 2
        for line in body_lines:
            ws_guidance.cell(row=current_row, column=1, value=f"•").font = regular_font
            c = ws_guidance.cell(row=current_row, column=2, value=_str(line))
            c.font = regular_font
            c.alignment = wrap_left
            c.border = thin_border
            ws_guidance.cell(row=current_row, column=1).border = thin_border
            ws_guidance.row_dimensions[current_row].height = min(max(26, len(_str(line)) // 70 * 15 + 26), 90)
            current_row += 1

    ws_guidance.column_dimensions["A"].width = 8
    ws_guidance.column_dimensions["B"].width = 110

    # SHEET 4: METADATA & ETHICS
    ws_meta = wb.create_sheet(title="Metadata & Ethics")
    ws_meta.sheet_properties.pageSetUpPr.fitToPage = True
    ws_meta.page_setup.fitToWidth = 1
    ws_meta.page_setup.fitToHeight = 1
    engine = report.get("engine_metadata", {}) or {}
    confidence = report.get("confidence", {}) or {}
    ws_meta["A1"] = "ENGINE METADATA & PROVENANCE (Rule 11)"
    ws_meta["A1"].font = section_font
    meta_rows = [
        ("Engine", _str(engine.get("engine", ""))),
        ("Engine version", _str(engine.get("version", ""))),
        ("Model", _str(engine.get("model", ""))),
        ("Generated (UTC)", _str(engine.get("timestamp_utc", ""))),
        ("Report ID", _str(report.get("report_id", ""))),
        ("Reading ID", _str(report.get("reading_id", ""))),
        ("Confidence", f"{_str(confidence.get('score', ''))}% ({_str(confidence.get('label', ''))})"),
        ("Review state", "Reviewed" if report.get("reviewed") else "Awaiting review"),
    ]
    for idx, (label, val) in enumerate(meta_rows, start=2):
        a = ws_meta.cell(row=idx, column=1, value=label)
        b = ws_meta.cell(row=idx, column=2, value=val)
        a.font = bold_font
        a.border = thin_border
        b.font = regular_font
        b.border = thin_border
        b.alignment = wrap_left

    ethics_row = len(meta_rows) + 4
    ws_meta[f"A{ethics_row}"] = "ETHICS & SAFETY NOTICE (§49.19)"
    ws_meta[f"A{ethics_row}"].font = bold_font
    ws_meta[f"A{ethics_row}"].fill = gold_fill
    ws_meta[f"A{ethics_row+1}"] = _str(report.get("disclaimer", ""))
    ws_meta[f"A{ethics_row+1}"].font = italic_font
    ws_meta[f"A{ethics_row+1}"].alignment = wrap_left
    ws_meta.row_dimensions[ethics_row + 1].height = 60
    ws_meta.merge_cells(f"A{ethics_row+1}:B{ethics_row+1}")
    ws_meta.column_dimensions["A"].width = 22
    ws_meta.column_dimensions["B"].width = 100

    for sheet in wb.worksheets:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
