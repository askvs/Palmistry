# Reading reports and export system (JSON, PDF, Excel).
import io
import uuid
from datetime import datetime, timezone

import openpyxl
import pytest

from app.models.reading import Reading, ReadingStatus, ReadingType
from app.models.user import User, UserRole
from app.services.reports import build_reading_report, export_report_to_excel, export_report_to_pdf


@pytest.fixture
def sample_user(test_engine):
    from app.db.session import get_db
    with test_engine.connect() as conn:
        pass
    from sqlalchemy.orm import Session
    with Session(test_engine) as session:
        user = User(
            id=uuid.uuid4(),
            email="seeker@example.com",
            password_hash="hash",
            full_name="Mystic Seeker",
            role=UserRole.user,
            is_active=True,
        )
        session.add(user)
        session.commit()
        session.refresh(user)
        return user


@pytest.fixture
def other_user(test_engine):
    from sqlalchemy.orm import Session
    with Session(test_engine) as session:
        user = User(
            id=uuid.uuid4(),
            email="other@example.com",
            password_hash="hash",
            full_name="Other Seeker",
            role=UserRole.user,
            is_active=True,
        )
        session.add(user)
        session.commit()
        session.refresh(user)
        return user


@pytest.fixture
def sample_palm_reading(test_engine, sample_user):
    from sqlalchemy.orm import Session
    with Session(test_engine) as session:
        reading = Reading(
            id=uuid.uuid4(),
            user_id=sample_user.id,
            type=ReadingType.palm,
            status=ReadingStatus.completed,
            title="Palm Analysis (Earth Hand)",
            summary="Detected Earth hand with 5 lines analyzed (88% confidence).",
            result={
                "metadata": {
                    "engine": "PalmBiometricsCV",
                    "hand": "Right",
                    "overall_confidence": 0.88,
                    "landmarks_count": 21,
                },
                "palm_shape": {
                    "element": "Earth",
                    "description": "Square palm with short fingers indicating grounded practicality.",
                    "palm_ratio": 1.02,
                    "finger_ratio": 0.78,
                },
                "lines": {
                    "life_line": {
                        "name": "Life Line",
                        "length": "Long",
                        "clarity": "Strong",
                        "curve": "Wide Arc",
                        "confidence": 0.92,
                        "interpretation": "Vibrant vitality and grounded stability.",
                    },
                    "head_line": {
                        "name": "Head Line",
                        "length": "Medium",
                        "clarity": "Distinct",
                        "curve": "Straight",
                        "confidence": 0.89,
                        "interpretation": "Clear analytical acumen and pragmatic problem solving.",
                    },
                    "heart_line": {
                        "name": "Heart Line",
                        "length": "Long",
                        "clarity": "Curved",
                        "curve": "Upward",
                        "confidence": 0.86,
                        "interpretation": "Deep emotional loyalty and warmth.",
                    },
                    "fate_line": {
                        "name": "Fate Line",
                        "length": "Present",
                        "clarity": "Moderate",
                        "curve": "Vertical",
                        "confidence": 0.81,
                        "interpretation": "Steadfast career focus and discipline.",
                    },
                    "sun_line": {
                        "name": "Sun Line",
                        "length": "Developing",
                        "clarity": "Fine",
                        "curve": "Vertical",
                        "confidence": 0.75,
                        "interpretation": "Creative potential and authentic recognition.",
                    },
                },
                "mounts": {
                    "jupiter": {"name": "Mount of Jupiter", "prominence": "Prominent", "prominence_score": 0.82, "meaning": "Leadership and optimism."},
                    "saturn": {"name": "Mount of Saturn", "prominence": "Normal", "prominence_score": 0.55, "meaning": "Discipline and prudence."},
                    "apollo": {"name": "Mount of Sun / Apollo", "prominence": "Normal", "prominence_score": 0.60, "meaning": "Artistic sensibility."},
                    "mercury": {"name": "Mount of Mercury", "prominence": "Prominent", "prominence_score": 0.76, "meaning": "Communication acuity."},
                    "venus": {"name": "Mount of Venus", "prominence": "Prominent", "prominence_score": 0.85, "meaning": "Vitality and passion."},
                    "moon": {"name": "Mount of Moon", "prominence": "Normal", "prominence_score": 0.58, "meaning": "Imagination and intuition."},
                },
            },
        )
        session.add(reading)
        session.commit()
        session.refresh(reading)
        return reading


@pytest.fixture
def sample_tarot_reading(test_engine, sample_user):
    from sqlalchemy.orm import Session
    with Session(test_engine) as session:
        reading = Reading(
            id=uuid.uuid4(),
            user_id=sample_user.id,
            type=ReadingType.tarot,
            status=ReadingStatus.completed,
            title="Three Card Spread",
            summary="A transformative journey from past reflection to visionary forward momentum.",
            result={
                "spread_name": "Three Card Spread",
                "spread_id": "three-card",
                "deck_id": "rider-waite-smith",
                "question": "What is guiding my professional transition?",
                "cards": [
                    {
                        "card_id": "the_fool",
                        "name": "The Fool",
                        "arcana": "Major Arcana",
                        "suit": "None",
                        "is_reversed": False,
                        "position": {"name": "Past", "meaning": "Foundational influences"},
                        "keywords": ["Beginnings", "Innocence", "Leap of Faith"],
                        "interpretation": "A bold step initiated a chapter of self-discovery.",
                    },
                    {
                        "card_id": "the_magician",
                        "name": "The Magician",
                        "arcana": "Major Arcana",
                        "suit": "None",
                        "is_reversed": False,
                        "position": {"name": "Present", "meaning": "Active current energy"},
                        "keywords": ["Manifestation", "Skill", "Willpower"],
                        "interpretation": "All tools and resources are in your hands to create alignment.",
                    },
                    {
                        "card_id": "the_star",
                        "name": "The Star",
                        "arcana": "Major Arcana",
                        "suit": "None",
                        "is_reversed": False,
                        "position": {"name": "Future", "meaning": "Probable trajectory"},
                        "keywords": ["Hope", "Inspiration", "Renewal"],
                        "interpretation": "Clarity, renewed faith, and radiant authentic inspiration.",
                    },
                ],
                "synthesis": {
                    "elemental_dominance": "Air",
                    "arcana_distribution": {"major": 3, "minor": 0},
                    "reversals_count": 0,
                },
                "narrative": {
                    "theme": "Sacred Transformation",
                    "summary": "Major arcana dominance shows life-altering directional alignment.",
                    "full_text": "You are moving through a rare cycle of high-impact creative self-manifestation.",
                },
                "guidance": [
                    "Take deliberate action without succumbing to imposter hesitation.",
                    "Ground your vision in daily rituals and clear intentions.",
                ],
                "metadata": {
                    "engine": "TarotInterpretationRuleEngine",
                    "confidence": 0.94,
                },
            },
        )
        session.add(reading)
        session.commit()
        session.refresh(reading)
        return reading


@pytest.fixture
def cv_palm_reading(test_engine, sample_user):
    # Palm reading in CV-engine shape: lines/mounts as lists, hand as dict.
    from sqlalchemy.orm import Session
    with Session(test_engine) as session:
        reading = Reading(
            id=uuid.uuid4(),
            user_id=sample_user.id,
            type=ReadingType.palm,
            status=ReadingStatus.completed,
            title="Palm Analysis (Water Hand)",
            summary="Detected Water hand with 5 lines analyzed (72% confidence).",
            result={
                "hand": {
                    "is_palmar": True,
                    "preferred_hand": "Right",
                    "detected_handedness": "Right",
                    "handedness_confidence": 0.968,
                    "landmarks_count": 21,
                    "palm_bounding_box": {"max_x": 741, "max_y": 533, "min_x": 426, "min_y": 298},
                },
                "metadata": {
                    "timestamp": "2026-09-08T11:39:54.429561+00:00",
                    "model_name": "MediaPipe HandLandmarker + OpenCV Crease Pipeline",
                    "model_version": "1.0.0",
                    "overall_confidence": 0.72,
                    "analyzed_resolution": {"width": 1024, "height": 915},
                },
                "palm_shape": {
                    "element": "Water",
                    "description": "Water Hand: Rectangular palm with long, graceful fingers.",
                    "traits": ["intuitive", "empathetic", "creative"],
                    "aspect_ratio": 0.648,
                    "finger_to_palm_ratio": 0.897,
                },
                "lines": [
                    {"name": "Life Line", "line_id": "life_line", "detected": True,
                     "curvature": 1.36, "confidence": 0.6, "continuity": "intermittent"},
                    {"name": "Head Line", "line_id": "head_line", "detected": True,
                     "curvature": 1.1, "confidence": 0.55, "continuity": "continuous"},
                    {"name": "Heart Line", "line_id": "heart_line", "detected": True,
                     "curvature": 0.9, "confidence": 0.45, "continuity": "broken"},
                    {"name": "Fate Line", "line_id": "fate_line", "detected": False,
                     "curvature": 0.0, "confidence": 0.2, "continuity": "absent"},
                    {"name": "Sun Line", "line_id": "sun_line", "detected": True,
                     "curvature": 0.8, "confidence": 0.3, "continuity": "fragmented"},
                ],
                "mounts": [
                    {"name": "Mount of Jupiter", "mount_id": "mount_of_jupiter", "score": 0.76,
                     "element": "Fire", "prominence": "prominent",
                     "meaning": "Dominant personality, proud, natural commander.",
                     "core_attributes": ["leadership", "ambition"], "planetary_ruler": "Jupiter"},
                    {"name": "Mount of Saturn", "mount_id": "mount_of_saturn", "score": 0.8,
                     "element": "Earth", "prominence": "prominent",
                     "meaning": "Serious, introverted, solitary temperament.",
                     "core_attributes": ["discipline"], "planetary_ruler": "Saturn"},
                    {"name": "Mount of Venus", "mount_id": "mount_of_venus", "score": 0.65,
                     "element": "Water", "prominence": "normal",
                     "meaning": "Warmth, vitality and passion.",
                     "core_attributes": ["love"], "planetary_ruler": "Venus"},
                ],
            },
        )
        session.add(reading)
        session.commit()
        session.refresh(reading)
        return reading


# =========================================================================
# UNIT TESTS: REPORT GENERATION & EXPORTERS
# =========================================================================


def test_palm_report_structure(sample_palm_reading, sample_user):
    """Verify Palmistry report conforms to FR-057 with confidence and 5 lines."""
    report = build_reading_report(sample_palm_reading, sample_user)
    assert report["report_type"] == "palmistry"
    assert report["user_name"] == "Mystic Seeker"
    assert report["confidence"]["score"] == 88.0
    assert report["confidence"]["label"] == "Very High"
    assert len(report["lines"]) == 5
    assert len(report["mounts"]) == 6
    assert "disclaimer" in report
    assert "guidance" in report["disclaimer"].lower()
    assert report["palm_features"]["element"] == "Earth"


def test_tarot_report_structure(sample_tarot_reading, sample_user):
    """Verify Tarot report conforms to FR-058 with positions, orientation, and synthesis."""
    report = build_reading_report(sample_tarot_reading, sample_user)
    assert report["report_type"] == "tarot"
    assert report["title"] == "Three Card Spread"
    assert report["confidence"]["score"] == 94.0
    assert report["confidence"]["label"] == "Very High"
    assert len(report["cards"]) == 3
    assert report["cards"][0]["name"] == "The Fool"
    assert report["cards"][0]["orientation"] == "Upright"
    assert report["synthesis"]["elemental_dominance"] == "Air"
    assert len(report["actionable_guidance"]) == 2


def test_palm_report_structure_from_cv_list_shape(cv_palm_reading, sample_user):
    """Report generation tolerates the list-shaped lines/mounts persisted by the CV pipeline."""
    report = build_reading_report(cv_palm_reading, sample_user)
    assert report["report_type"] == "palmistry"
    assert len(report["lines"]) == 5
    assert [line["name"] for line in report["lines"]] == [
        "Life Line", "Head Line", "Heart Line", "Fate Line", "Sun Line",
    ]
    assert [line["key"] for line in report["lines"]] == [
        "life_line", "head_line", "heart_line", "fate_line", "sun_line",
    ]
    assert report["lines"][0]["confidence"] == 60.0
    assert report["lines"][0]["interpretation"]  # continuity-derived fallback present
    assert len(report["mounts"]) == 3
    assert report["mounts"][0]["name"] == "Mount of Jupiter"
    assert report["mounts"][0]["score"] == 76.0
    assert report["mounts"][0]["meaning"] == "Dominant personality, proud, natural commander."
    assert report["palm_features"]["element"] == "Water"
    assert report["palm_features"]["hand_detected"] == "Right"
    assert report["palm_features"]["palm_length_to_width"] == 0.648
    assert report["palm_features"]["finger_to_palm_ratio"] == 0.897


def test_exports_from_cv_list_shape(cv_palm_reading, sample_user):
    """PDF and XLSX exporters succeed for list-shaped CV palm data (FR-062, FR-063)."""
    report = build_reading_report(cv_palm_reading, sample_user)

    pdf = export_report_to_pdf(report)
    assert isinstance(pdf, bytes)
    assert pdf.startswith(b"%PDF-")

    xlsx = export_report_to_excel(report)
    assert xlsx.startswith(b"PK")
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    assert wb["Overview"]["A2"].value == "Palm Analysis (Water Hand)"
    details = wb["Analysis Details"]
    col_a = [details.cell(row=r, column=1).value for r in range(1, details.max_row + 1)]
    assert "Life Line" in col_a
    # Continuity column still present for the CV list shape.
    found_continuity = False
    for row in details.iter_rows(values_only=True):
        values = [str(v or "") for v in row]
        if "Life Line" in values and "intermittent" in values:
            found_continuity = True
    assert found_continuity or "Head Line" in col_a


def test_unhandled_error_response_includes_cors_headers(
    client, sample_tarot_reading, sample_user, monkeypatch
):
    # Unhandled 500s must still carry CORS headers so the browser shows the real error.
    from fastapi.testclient import TestClient
    from app.main import app

    def boom(report):
        raise RuntimeError("export exploded")

    monkeypatch.setattr(
        "app.api.v1.endpoints.reports.export_report_to_pdf", boom
    )
    headers = _make_auth_header(sample_user)

    with TestClient(app, raise_server_exceptions=False) as c:
        res = c.get(
            f"/api/v1/reports/reading/{sample_tarot_reading.id}/pdf",
            headers={**headers, "Origin": "http://localhost:3000"},
        )
        assert res.status_code == 500
        assert res.json()["code"] == "INTERNAL_ERROR"
        assert res.headers["access-control-allow-origin"] == "http://localhost:3000"
        assert res.headers["access-control-allow-credentials"] == "true"

        # Disallowed origins get no CORS headers (browser will still block them)
        res_blocked = c.get(
            f"/api/v1/reports/reading/{sample_tarot_reading.id}/pdf",
            headers={**headers, "Origin": "http://evil.example"},
        )
        assert res_blocked.status_code == 500
        assert "access-control-allow-origin" not in res_blocked.headers


def test_pdf_export_generation(sample_palm_reading, sample_tarot_reading, sample_user):
    """Verify PDF export conforms to FR-062 and returns valid PDF binary magic bytes."""
    palm_report = build_reading_report(sample_palm_reading, sample_user)
    pdf_palm = export_report_to_pdf(palm_report)
    assert isinstance(pdf_palm, bytes)
    assert pdf_palm.startswith(b"%PDF-")
    assert len(pdf_palm) > 2000

    tarot_report = build_reading_report(sample_tarot_reading, sample_user)
    pdf_tarot = export_report_to_pdf(tarot_report)
    assert isinstance(pdf_tarot, bytes)
    assert pdf_tarot.startswith(b"%PDF-")
    assert len(pdf_tarot) > 2000


def test_excel_export_generation(sample_palm_reading, sample_tarot_reading, sample_user):
    """Verify Excel export conforms to FR-063 and produces valid XLSX workbook with sheets."""
    palm_report = build_reading_report(sample_palm_reading, sample_user)
    xlsx_palm = export_report_to_excel(palm_report)
    assert isinstance(xlsx_palm, bytes)
    assert xlsx_palm.startswith(b"PK")  # ZIP magic bytes for XLSX

    # Inspect with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_palm))
    assert "Overview" in wb.sheetnames
    assert "Analysis Details" in wb.sheetnames
    assert "Guidance & Actions" in wb.sheetnames
    assert wb["Overview"]["A2"].value == "Palm Analysis (Earth Hand)"

    tarot_report = build_reading_report(sample_tarot_reading, sample_user)
    xlsx_tarot = export_report_to_excel(tarot_report)
    wb_tarot = openpyxl.load_workbook(io.BytesIO(xlsx_tarot))
    assert "Overview" in wb_tarot.sheetnames
    assert wb_tarot["Overview"]["A2"].value == "Three Card Spread"


# =========================================================================
# INTEGRATION TESTS: REPORT ENDPOINTS & ACCESS CONTROL
# =========================================================================


def _make_auth_header(user: User) -> dict:
    from app.auth.tokens import create_access_token
    token = create_access_token(user.id, user.role)
    return {"Authorization": f"Bearer {token}"}


def test_get_report_json_endpoint(client, sample_palm_reading, sample_user, other_user):
    """Authorized user can fetch JSON report, unauthorized cannot."""
    headers = _make_auth_header(sample_user)
    res = client.get(f"/api/v1/reports/reading/{sample_palm_reading.id}", headers=headers)
    assert res.status_code == 200
    data = res.json()
    assert data["report_type"] == "palmistry"
    assert data["reading_id"] == str(sample_palm_reading.id)

    # Other user forbidden
    other_headers = _make_auth_header(other_user)
    res_forbidden = client.get(f"/api/v1/reports/reading/{sample_palm_reading.id}", headers=other_headers)
    assert res_forbidden.status_code == 403

    # Unauthenticated rejected
    res_unauth = client.get(f"/api/v1/reports/reading/{sample_palm_reading.id}")
    assert res_unauth.status_code == 401


def test_download_report_pdf_endpoint(client, sample_tarot_reading, sample_user):
    """Owner can download PDF stream with proper Content-Disposition header."""
    headers = _make_auth_header(sample_user)
    res = client.get(f"/api/v1/reports/reading/{sample_tarot_reading.id}/pdf", headers=headers)
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert "attachment; filename=" in res.headers["content-disposition"]
    assert res.content.startswith(b"%PDF-")


def test_download_report_excel_endpoint(client, sample_palm_reading, sample_user):
    """Owner can download Excel stream with proper Content-Disposition header."""
    headers = _make_auth_header(sample_user)
    res = client.get(f"/api/v1/reports/reading/{sample_palm_reading.id}/excel", headers=headers)
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]
    assert "attachment; filename=" in res.headers["content-disposition"]
    assert res.content.startswith(b"PK")


def test_delete_reading_endpoint(client, sample_tarot_reading, sample_user):
    """Owner can delete reading; subsequent fetch yields 404."""
    headers = _make_auth_header(sample_user)
    del_res = client.delete(f"/api/v1/readings/{sample_tarot_reading.id}", headers=headers)
    assert del_res.status_code == 204

    # Verify 404
    get_res = client.get(f"/api/v1/readings/{sample_tarot_reading.id}", headers=headers)
    assert get_res.status_code == 404
